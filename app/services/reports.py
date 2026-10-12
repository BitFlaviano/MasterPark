"""Serviço de relatórios: geração em CSV, Excel e PDF."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from app.domain.models import format_money, format_pt_datetime

REPORTS = {
    "entradas": {"title": "Entradas", "columns": ["Ticket", "Placa", "Entrada", "Operador", "Terminal", "Vaga"]},
    "saidas": {"title": "Saídas", "columns": ["Ticket", "Placa", "Entrada", "Saída", "Minutos", "Valor"]},
    "presentes": {"title": "Veículos Presentes", "columns": ["Ticket", "Placa", "Entrada", "Vaga", "Tempo"]},
    "faturamento": {"title": "Faturamento", "columns": ["Ticket", "Placa", "Entrada", "Saída", "Minutos", "Base", "Desconto", "Final", "Status"]},
    "mensalistas": {"title": "Mensalistas", "columns": ["Cliente", "CPF/CNPJ", "Telefone", "Plano", "Valor", "Status"]},
    "inadimplentes": {"title": "Inadimplentes", "columns": ["Cliente", "Referência", "Vencimento", "Valor"]},
    "convenios": {"title": "Convênios", "columns": ["Convênio", "Tipo Desconto", "Valor", "Ativo"]},
    "descontos": {"title": "Descontos", "columns": ["Ticket", "Tipo", "Valor", "Motivo", "Data"]},
    "cancelamentos": {"title": "Cancelamentos", "columns": ["Ticket", "Placa", "Cancelado em", "Motivo"]},
    "operadores": {"title": "Operadores", "columns": ["Operador", "Tickets", "Faturamento"]},
    "auditoria": {"title": "Auditoria", "columns": ["Data", "Usuário", "Ação", "Entidade", "Registro", "Detalhe"]},
    "caixa": {"title": "Movimentações de Caixa", "columns": ["Registro", "Data", "Operador", "Tipo", "Valor", "Descrição"]},
    "pagamentos": {"title": "Pagamentos", "columns": ["Ticket", "Método", "Status", "Valor", "Recebido", "Recebido em"]},
    "tarifas": {"title": "Tarifas", "columns": ["Nome", "Categoria", "Período", "Regra", "Valor"]},
    "ocupacao": {"title": "Ocupação", "columns": ["Vaga", "Setor", "Tipo", "Estado"]},
    "picosh": {"title": "Horários de Pico", "columns": ["Hora", "Entradas", "Saídas", "Faturamento"]},
}


class ReportService:
    def __init__(self, db, lot=None):
        self.db = db
        self.lot = lot

    # ------------------------------------------------------------------
    # Construção dos dados
    # ------------------------------------------------------------------
    def build(self, report: str, start: str = "", end: str = "",
              term: str = "") -> tuple[str, list[str], list[list]]:
        if report not in REPORTS:
            raise ValueError(f"Relatório desconhecido: {report}")
        title, cols = REPORTS[report]["title"], REPORTS[report]["columns"]
        rows = self._rows(report, start, end, term)
        return title, cols, rows

    def _rows(self, report: str, start: str, end: str, term: str) -> list[list]:
        db = self.db
        cond_time = ""
        params: list = []
        if start and end:
            key = "entry_at" if report in ("entradas", "presentes") else "exit_at"
            cond_time = f" AND t.{key}>=? AND t.{key}<=?"
            params += [start, end]

        if report == "entradas":
            rows = db.query(
                "SELECT t.number, t.plate, t.entry_at, u.name, t.terminal_entry, s.code"
                " FROM tickets t LEFT JOIN users u ON u.id=t.entry_user_id"
                " LEFT JOIN parking_spaces s ON s.id=t.parking_space_id"
                f" WHERE 1=1 {cond_time} ORDER BY t.id DESC", tuple(params))
            return [[r["number"], r["plate"], format_pt_datetime(r["entry_at"]),
                     r["name"] or "-", r["terminal_entry"], r["code"] or "-"] for r in rows]

        if report == "saidas":
            rows = db.query(
                "SELECT t.number, t.plate, t.entry_at, t.exit_at, t.total_minutes, t.final_value"
                " FROM tickets t WHERE t.status IN ('FINALIZADO','PERDIDO')"
                f" {cond_time} ORDER BY t.exit_at DESC", tuple(params))
            return [[r["number"], r["plate"], format_pt_datetime(r["entry_at"]),
                     format_pt_datetime(r["exit_at"]), r["total_minutes"],
                     format_money(r["final_value"])] for r in rows]

        if report == "presentes":
            rows = db.query(
                "SELECT t.number, t.plate, t.entry_at, s.code, t.vehicle_category"
                " FROM tickets t LEFT JOIN parking_spaces s ON s.id=t.parking_space_id"
                " WHERE t.status='ATIVO' ORDER BY t.entry_at", ())
            out = []
            for r in rows:
                mins = _elapsed_minutes(r["entry_at"])
                h, m = divmod(mins, 60)
                out.append([r["number"], r["plate"], format_pt_datetime(r["entry_at"]),
                            r["code"] or "-", f"{h}h{m:02d}"])
            return out

        if report == "faturamento":
            rows = db.query(
                "SELECT t.number, t.plate, t.entry_at, t.exit_at, t.total_minutes,"
                " t.base_value, t.discount_value, t.final_value, t.status FROM tickets t"
                f" WHERE 1=1 {cond_time} ORDER BY t.exit_at DESC", tuple(params))
            return [[r["number"], r["plate"], format_pt_datetime(r["entry_at"]),
                     format_pt_datetime(r["exit_at"]), r["total_minutes"] or "-",
                     format_money(r["base_value"]), format_money(r["discount_value"]),
                     format_money(r["final_value"]), r["status"]] for r in rows]

        if report == "mensalistas":
            rows = db.query(
                "SELECT c.name, COALESCE(c.cpf, c.cnpj), c.phone, m.plan_name,"
                " m.monthly_value, m.status FROM monthly_customers m"
                " JOIN customers c ON c.id=m.customer_id ORDER BY c.name")
            return [[r["name"], r["cpf"] or r["cnpj"] or "-", r["phone"] or "-",
                     r["plan_name"], format_money(r["monthly_value"]), r["status"]] for r in rows]

        if report == "inadimplentes":
            rows = db.query(
                "SELECT c.name, mp.reference, mp.due_date, mp.value FROM monthly_payments mp"
                " JOIN monthly_customers m ON m.id=mp.monthly_customer_id"
                " JOIN customers c ON c.id=m.customer_id"
                " WHERE mp.status='PENDENTE' AND mp.due_date < date('now')"
                " ORDER BY mp.due_date")
            return [[r["name"], r["reference"], format_pt_datetime(r["due_date"]),
                     format_money(r["value"])] for r in rows]

        if report == "convenios":
            rows = db.query("SELECT * FROM agreements ORDER BY name")
            return [[r["name"], r["discount_type"], format_money(r["discount_value"])
                     if r["discount_type"] != "HORAS" else f"{r['discount_hours']}h",
                     "Sim" if r["active"] else "Não"] for r in rows]

        if report == "descontos":
            rows = db.query(
                "SELECT t.number, d.kind, d.value, d.reason, d.created_at FROM discounts d"
                " LEFT JOIN tickets t ON t.id=d.ticket_id"
                f" WHERE d.created_at>=? AND d.created_at<=? ORDER BY d.id DESC",
                (start, end)) if start and end else db.query(
                "SELECT t.number, d.kind, d.value, d.reason, d.created_at FROM discounts d"
                " LEFT JOIN tickets t ON t.id=d.ticket_id ORDER BY d.id DESC")
            return [[r["number"], r["kind"], format_money(r["value"]), r["reason"],
                     format_pt_datetime(r["created_at"])] for r in rows]

        if report == "cancelamentos":
            rows = db.query(
                "SELECT t.number, t.plate, t.exit_at, t.cancel_reason FROM tickets t"
                " WHERE t.status='CANCELADO' ORDER BY t.exit_at DESC")
            return [[r["number"], r["plate"], format_pt_datetime(r["exit_at"]),
                     r["cancel_reason"]] for r in rows]

        if report == "operadores":
            rows = db.query(
                "SELECT u.name, COUNT(*) n, COALESCE(SUM(t.final_value),0) v FROM tickets t"
                " JOIN users u ON u.id=t.exit_user_id"
                " WHERE t.status IN ('FINALIZADO','PERDIDO')"
                f" {cond_time} GROUP BY u.name ORDER BY v DESC", tuple(params))
            return [[r["name"], r["n"], format_money(r["v"])] for r in rows]

        if report == "auditoria":
            rows = db.query(
                "SELECT created_at, username, action, entity, entity_id,"
                " COALESCE(old_value,'')||' → '||COALESCE(new_value,'') detail"
                " FROM audit_logs ORDER BY id DESC LIMIT 1000")
            return [[format_pt_datetime(r["created_at"]), r["username"], r["action"],
                     r["entity"], r["entity_id"], r["detail"]] for r in rows]

        if report == "caixa":
            rows = db.query(
                "SELECT cm.cash_register_id, cm.created_at, u.name, cm.kind, cm.value,"
                " cm.description FROM cash_movements cm LEFT JOIN users u ON u.id=cm.operator_id"
                " ORDER BY cm.id DESC")
            return [[r["cash_register_id"], format_pt_datetime(r["created_at"]), r["name"] or "-",
                     r["kind"], format_money(r["value"]), r["description"]] for r in rows]

        if report == "pagamentos":
            rows = db.query(
                "SELECT t.number, pm.name, p.status, p.value, p.received_value, p.received_at"
                " FROM payments p LEFT JOIN tickets t ON t.id=p.ticket_id"
                " JOIN payment_methods pm ON pm.code=p.payment_method"
                " ORDER BY p.id DESC LIMIT 1000")
            return [[r["number"] or "-", r["name"], r["status"], format_money(r["value"]),
                     format_money(r["received_value"]),
                     format_pt_datetime(r["received_at"])] for r in rows]

        if report == "tarifas":
            rows = db.query(
                "SELECT t.name, t.vehicle_category, t.period_type, r.rule_type, r.value"
                " FROM tariffs t LEFT JOIN tariff_rules r ON r.tariff_id=t.id"
                " ORDER BY t.name, r.minutes_from")
            return [[r["name"], r["vehicle_category"], r["period_type"], r["rule_type"],
                     format_money(r["value"]) if r["value"] else "-"] for r in rows]

        if report == "ocupacao":
            rows = db.query(
                "SELECT s.code, sec.name, s.space_type, s.state FROM parking_spaces s"
                " LEFT JOIN parking_sectors sec ON sec.id=s.sector_id"
                " WHERE s.active=1 ORDER BY s.code")
            return [[r["code"], r["name"] or "-", r["space_type"], r["state"]] for r in rows]

        if report == "picosh":
            rows = db.query(
                "SELECT strftime('%H', entry_at) h, COUNT(*) c FROM tickets"
                " WHERE entry_at>=? AND entry_at<=? GROUP BY h ORDER BY c DESC",
                (start, end)) if start and end else []
            return [[f"{r['h']}h", r["c"], "", ""] for r in rows]

        return []

    # ------------------------------------------------------------------
    # Exportação
    # ------------------------------------------------------------------
    def to_csv(self, title: str, cols: list[str], rows: list[list]) -> str:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(cols)
        writer.writerows(rows)
        return buf.getvalue()

    def save_csv(self, report: str, start: str = "", end: str = "",
                 out_dir: str = "") -> str:
        title, cols, rows = self.build(report, start, end)
        import os
        out = out_dir or "relatorios"
        os.makedirs(out, exist_ok=True)
        path = os.path.join(out, f"{report}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(cols)
            writer.writerows(rows)
        return path

    def to_excel(self, title: str, cols: list[str], rows: list[list]) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        header = ws.cell(1, 1, value=title)
        header.font = Font(bold=True, size=14)
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(3, c, value=col)
            cell.font = Font(bold=True)
        for r, row in enumerate(rows, start=4):
            for c, val in enumerate(row, start=1):
                ws.cell(r, c, value=val)
        ws.freeze_panes = "A4"
        for idx, _ in enumerate(cols, start=1):
            ws.column_dimensions[chr(64 + idx if idx <= 26 else 65)].width = 18
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def to_pdf(self, title: str, cols: list[str], rows: list[list],
               subtitle: str = "") -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                        TableStyle)
        from reportlab.lib.styles import getSampleStyleSheet

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                leftMargin=12 * mm, rightMargin=12 * mm,
                                topMargin=12 * mm, bottomMargin=12 * mm)
        styles = getSampleStyleSheet()
        story = [Paragraph(title, styles["Title"])]
        if subtitle:
            story.append(Paragraph(subtitle, styles["Small"]))
        story.append(Spacer(1, 6 * mm))
        data = [cols] + rows
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4C6FFF")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4FF")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
        doc.build(story)
        return buf.getvalue()


def _elapsed_minutes(entry_iso: str) -> int:
    from datetime import datetime
    try:
        entry = datetime.fromisoformat(entry_iso.replace("Z", "+00:00"))
        return int((datetime.now(entry.tzinfo) - entry).total_seconds() // 60)
    except Exception:
        return 0